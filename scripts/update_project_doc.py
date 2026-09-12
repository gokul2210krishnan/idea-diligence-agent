import os
import shutil
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def main():
    docx_path = "Idea Diligence Agent - Project.docx"
    backup_path = "Idea Diligence Agent - Project.docx.bak"

    if not os.path.exists(backup_path):
        shutil.copyfile(docx_path, backup_path)
        print(f"Created backup at {backup_path}")

    doc = docx.Document(docx_path)

    # Find the target insertion point: right before "Gokul's Points"
    gokul_p_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Gokul's Points":
            gokul_p_idx = i
            break

    if gokul_p_idx is None:
        raise ValueError("Could not find 'Gokul\\'s Points' paragraph in document")

    target_p = doc.paragraphs[gokul_p_idx]

    # Helper function to insert before target_p
    def add_p(text="", style='normal', bold=False, italic=False, space_after=6):
        p = target_p.insert_paragraph_before()
        if text:
            run = p.add_run(text)
            run.bold = bold
            run.italic = italic
        try:
            p.style = style
        except Exception:
            p.style = 'normal'
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def add_h(text, level=1):
        style_name = f'Heading {level}'
        p = target_p.insert_paragraph_before()
        run = p.add_run(text)
        run.bold = True
        try:
            p.style = style_name
        except Exception:
            p.style = 'normal'
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        return p

    def add_bullet(text, bold_prefix="", space_after=4):
        p = target_p.insert_paragraph_before()
        try:
            p.style = 'normal'
        except Exception:
            pass
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(space_after)
        bullet_run = p.add_run("• ")
        bullet_run.bold = True
        if bold_prefix:
            prefix_run = p.add_run(bold_prefix)
            prefix_run.bold = True
        p.add_run(text)
        return p

    # Insert Edmund's Points
    add_p("Below is Edmund's refined project specification, architecture blueprint, and development guide. This synthesizes Gokul's feedback into an actionable, development-ready plan designed for all team members—including junior developers—to read, understand, and start building immediately.", italic=True, space_after=12)

    add_h("1. Executive Summary & Response to Gokul's Architectural Feedback", level=1)
    add_p("I completely agree with Gokul's key architectural insights. His review pushed our thinking from a rigid, linear script into a robust, autonomous agent system. Specifically:")
    add_bullet(" Instead of treating state as an ad-hoc database or simple chat history, the investigation centers on a single canonical data structure—the DiligenceState. Every specialist agent reads from it and contributes structured updates.", bold_prefix="Canonical Structured State: ")
    add_bullet(" Instead of running Problem → Competition → Economics in a single fixed line, the Orchestrator runs an adaptive loop. At each step, it identifies the single highest-value unknown and dispatches the appropriate specialist to resolve it.", bold_prefix="Adaptive Research Loop: ")
    add_bullet(" Each idea is investigated in an isolated session (ResearchSession) with a fresh DiligenceState, generating a persistent final report and cleanly disposing active memory so that batch processing (e.g. 10 ideas) is trivial and state leaks never happen.", bold_prefix="Session Isolation: ")
    add_bullet(" We do not treat all web data equally. Findings are strictly classified into FACT, ASSUMPTION, INFERENCE, and UNKNOWN, giving our orchestrator and users clear confidence boundaries.", bold_prefix="Evidence Classification: ")
    add_bullet(" We will not build a 40-field schema before testing real queries. We start with minimal core fields, validate them end-to-end with the Strands SDK and Amazon Bedrock, and let real research workflows guide model expansion.", bold_prefix="Lean Development Philosophy: ")

    add_h("2. Problem Statement & User Experience", level=1)
    add_p("Why We Are Building This:")
    add_bullet("Founders, indie hackers, and product managers waste dozens of hours manually Googling competitors, reading pricing tiers, guessing addressable market sizes, and compiling spreadsheets.")
    add_bullet("Founders are prone to severe confirmation bias—searching only for evidence that validates their enthusiasm while missing fatal red flags or strong incumbent moats.")
    add_bullet("Most people either build for months without basic market diligence, or prematurely kill viable concepts due to lack of competitive clarity.")
    
    add_p("The User Experience (Zero Babysitting):")
    add_bullet(" The user enters a plain-English product concept (e.g., 'An AI tool that tracks inventory waste for independent restaurants').", bold_prefix="1. Submission: ")
    add_bullet(" The agent triggers an autonomous investigation in the background. The user does not need to answer 20 tedious prompt questions.", bold_prefix="2. Autonomous Research: ")
    add_bullet(" The agent searches the web, parses competitor pricing tables, identifies target buyer segments, and evaluates unit economics.", bold_prefix="3. Multi-Turn Diligence: ")
    add_bullet(" The agent stress-tests the founder's assumptions against hard market evidence and competing alternatives.", bold_prefix="4. Assumption Challenge: ")
    add_bullet(" The agent delivers a clear, evidence-backed verdict (GO, MODIFY, or KILL) alongside a structured diligence dossier.", bold_prefix="5. Decision Dossier: ")

    add_h("3. System Architecture & The Diligence Loop", level=1)
    add_p("The system architecture consists of four primary components operating over an isolated session:")

    add_h("Component A: The Canonical Research State (DiligenceState)", level=2)
    add_p("DiligenceState is our single source of truth for an investigation. It contains:")
    add_bullet("Raw text description and initial user inputs.", bold_prefix="• idea: ")
    add_bullet("Analysis of customer pain points, severity, and existing manual workarounds.", bold_prefix="• problem: ")
    add_bullet("List of target personas, buyer profiles, and willingness-to-pay indicators.", bold_prefix="• customers: ")
    add_bullet("Direct competitors, indirect substitutes, pricing tiers, and feature gaps.", bold_prefix="• competitors: ")
    add_bullet("Pricing model options, unit margins, and estimated market size (TAM).", bold_prefix="• economics: ")
    add_bullet("Audit log of all gathered evidence categorized by epistemic type.", bold_prefix="• evidence: ")
    add_bullet("List of remaining high-priority uncertainties that require research.", bold_prefix="• unknowns: ")
    add_bullet("Step-by-step record of actions, agent invocations, and queries.", bold_prefix="• research_history: ")
    add_bullet("Final decision (GO, MODIFY, KILL) with confidence score and reasoning.", bold_prefix="• verdict: ")

    add_h("Evidence Classification Hierarchy", level=2)
    add_bullet("Verified, sourced empirical facts (e.g., 'Competitor X charges $49/mo on their public pricing page'). Must include a URL or citation.", bold_prefix="FACT: ")
    add_bullet("Unverified hypotheses or claims made by the user (e.g., 'Small restaurants will readily adopt an iPad app for waste tracking').", bold_prefix="ASSUMPTION: ")
    add_bullet("Logical deductions formed by connecting verified facts (e.g., 'Because 3 competitors offer tiered pricing above $50, the market has willingness to pay in this range').", bold_prefix="INFERENCE: ")
    add_bullet("Critical missing data that prevents confident decision-making (e.g., 'Average customer churn rate in this vertical').", bold_prefix="UNKNOWN: ")

    add_h("Component B: The Adaptive Orchestrator", level=2)
    add_p("The Orchestrator acts as the lead investigator. Instead of blindly executing agents in sequence, it follows an adaptive loop:")
    add_bullet("1. Inspect current DiligenceState and extract the list of critical UNKNOWNS.")
    add_bullet("2. Prioritize the single most valuable unknown to resolve next.")
    add_bullet("3. Dispatch the specialist agent best suited to answer it (Problem, Competition, or Economics).")
    add_bullet("4. Ingest the specialist's structured findings and merge them into DiligenceState.")
    add_bullet("5. Evaluate stopping conditions: If evidence is sufficient to make a confident recommendation OR maximum cycle limit (e.g., 5 cycles) is reached, trigger the Verdict Engine; otherwise, repeat.")

    add_h("Component C: The Three Specialist Research Agents", level=2)
    add_bullet("Investigates who experiences the problem, how acute the pain is, what manual workarounds exist today, and whether active demand signals exist.", bold_prefix="1. Problem Analysis Agent: ")
    add_bullet("Scrapes search results for direct competitors and indirect alternatives, finds published pricing plans, uncovers customer complaints on review sites, and pinpoints market whitespace.", bold_prefix="2. Competition Analysis Agent: ")
    add_bullet("Evaluates pricing models (subscription, usage, tiered), customer acquisition cost (CAC) vs. lifetime value (LTV) risks, and total addressable market feasibility.", bold_prefix="3. Economics Analysis Agent: ")

    add_h("Component D: Research Tools", level=2)
    add_bullet("Live web search tool (DuckDuckGo search) enabling agents to find live websites, press releases, and reviews without requiring paid API keys.", bold_prefix="search_web: ")
    add_bullet("HTTP fetch and markdown extractor tool that retrieves webpage content and strips noise for LLM consumption.", bold_prefix="read_webpage: ")
    add_bullet("A structured recording tool that forces agents to tag findings as FACT, ASSUMPTION, INFERENCE, or UNKNOWN with citation URLs.", bold_prefix="save_finding: ")

    add_h("4. The Decision Verdict Framework", level=1)
    add_p("At the end of an investigation, the agent synthesizes all evidence and renders one of three definitive verdicts:")
    add_bullet("Strong problem-solution fit, verified customer pain, identifiable competitive gap, defensible unit economics, and validated key assumptions. The recommendation is to proceed to MVP development.", bold_prefix="GO: ")
    add_bullet("The overarching problem is real, but fatal obstacles exist in the current formulation (e.g., crowded direct competition, poor pricing power, or wrong customer segment). The agent outputs exact pivot recommendations (e.g., 'Pivot from independent cafes to multi-location franchises').", bold_prefix="MODIFY: ")
    add_bullet("The idea faces insurmountable headwinds: entrenched free incumbents, negligible willingness to pay, severe regulatory hurdles, or an unverified core problem. The agent provides clear evidence saving the founder months of wasted effort.", bold_prefix="KILL: ")

    add_h("5. Technology Stack & AWS Alignment", level=1)
    add_p("To maximize our technical score in the AWS Strands Agents SDK Hackathon, our implementation specifically leverages:")
    add_bullet("Strands Agents SDK (Python) utilizing both the Agents-as-Tools pattern for rapid specialist delegation and the Graph pattern for cyclic feedback loops with bounded execution.", bold_prefix="Agent Framework: ")
    add_bullet("Amazon Bedrock running Anthropic Claude Sonnet, giving us industry-leading reasoning, accurate structured JSON outputs, and fast execution within AWS.", bold_prefix="Foundation Model: ")
    add_bullet("Pydantic v2 data models for rigorous schema enforcement, type safety, and clean JSON serialization across all agent interfaces.", bold_prefix="Data Modeling: ")
    add_bullet("Clean Python CLI with rich terminal formatting for local execution, with modular endpoints designed for rapid web UI or AWS AgentCore deployment.", bold_prefix="Interface: ")

    add_h("6. Development Roadmap for the Team & Junior Developers", level=1)
    add_p("Here is our step-by-step build sequence for the hackathon sprint:")
    add_bullet("Set up project skeleton, dependencies, Pydantic data models (models.py), custom research tools (tools.py), and test Bedrock connectivity. [COMPLETED]", bold_prefix="Phase 1 - Walking Skeleton: ")
    add_bullet("Implement the specialist agents (Problem, Competition, Economics) with targeted system prompts and wire the Orchestrator loop using the agents-as-tools coordinator.", bold_prefix="Phase 2 - Multi-Agent Loop: ")
    add_bullet("Refine the evidence classifier (FACT/ASSUMPTION/INFERENCE/UNKNOWN), add confidence scoring, and build the GO/MODIFY/KILL verdict generator.", bold_prefix="Phase 3 - Verdict & Synthesis Engine: ")
    add_bullet("Create interactive CLI/UI demo, record the 5-minute hackathon walkthrough video, write README documentation, and submit to Devpost.", bold_prefix="Phase 4 - Demo, Polish & Submission: ")

    add_h("7. Golden Rules for Developers Working on This Codebase", level=1)
    add_bullet("Do not add fields to DiligenceState until a real research run proves that an agent needs to produce or consume that data.", bold_prefix="1. Workflow Shapes the Model: ")
    add_bullet("Never allow an agent to label a statement as a FACT without an accompanying source URL or reference.", bold_prefix="2. Mandatory Provenance: ")
    add_bullet("Each specialist agent must focus strictly on its designated domain. Problem does not do pricing math; Economics does not do competitor discovery.", bold_prefix="3. Separation of Concerns: ")
    add_bullet("Every loop must have a hard execution cap (e.g. max 5 cycles) so the agent never hangs or consumes unbounded tokens.", bold_prefix="4. Bounded Execution: ")

    add_p("----------------------------------------------------------------------", space_after=12)

    doc.save(docx_path)
    print(f"Successfully saved updated {docx_path}")

if __name__ == "__main__":
    main()
